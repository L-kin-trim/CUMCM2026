"""验证修订Excel保留原有数字，检查DOCX结构并渲染Word导出PDF。"""
from pathlib import Path
import json,hashlib,sys,zipfile
import openpyxl
from lxml import etree
ROOT=Path(__file__).resolve().parents[1];REVIEW=ROOT.parent;PROJECT=REVIEW.parent
def sheets():
    reports=[]
    for group,name in [('问题1','result1'),('问题2','result2'),('问题3','result3'),('问题4','result4-2'),('问题4','result4-3')]:
        old=openpyxl.load_workbook(PROJECT/(name+'_final.xlsx'),data_only=True)
        dest=REVIEW/'06_错误说明与纠正'/group/(name+'_修订.xlsx')
        new=openpyxl.load_workbook(dest,data_only=True)
        changes=[];maxerr=0.;count=0
        assert new.sheetnames==old.sheetnames+['费用核算']
        for sn in old.sheetnames:
            a,b=old[sn],new[sn]
            assert (a.max_row,a.max_column)==(b.max_row,b.max_column)
            for row in a:
                for c in row:
                    x,y=c.value,b[c.coordinate].value
                    if isinstance(x,(int,float)):
                        assert isinstance(y,(int,float)),(name,sn,c.coordinate,x,y)
                        maxerr=max(maxerr,abs(x-y));count+=1
                    elif x!=y:
                        if hasattr(x,'isoformat') and hasattr(y,'isoformat') and x==y:continue
                        changes.append([sn,c.coordinate,str(x),str(y)])
        allowed=[['计划购电量','A145' if name=='result1' else 'EO1']]
        if '调整购电量' in old.sheetnames:allowed.append(['调整购电量','EO1'])
        assert [[r[0],r[1]] for r in changes]==allowed,(name,changes)
        assert maxerr<1e-7
        if name!='result1':
            s=new['费用核算']
            err=max(abs(s.cell(r,11).value-(s.cell(r,5).value+s.cell(r,6).value+s.cell(r,7).value-s.cell(r,8).value+s.cell(r,10).value)) for r in range(2,336))
            assert err<1e-6,err
        reports.append({'file':dest.name,'numeric_cells_compared':count,'max_numeric_difference':maxerr,'text_changes':changes})
    (REVIEW/'08_验证记录/修订Excel数值保留验证.json').write_text(json.dumps(reports,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(reports,ensure_ascii=False,indent=2))
def docx():
    p=REVIEW/'05_论文/微网电力调控论文_审查修订版.docx'
    with zipfile.ZipFile(p) as z:
        x=etree.fromstring(z.read('word/document.xml'))
        ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math','w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        report={'native_equations':len(x.findall('.//m:oMath',ns)),'display_equations':len(x.findall('.//m:oMathPara',ns)),
                'nested_math_paragraphs':len(x.findall('.//m:oMathPara/m:oMathPara',ns)),
                'tables':len(x.findall('.//w:tbl',ns)),'page_margins':[dict(v.attrib) for v in x.findall('.//w:pgMar',ns)]}
        assert report['native_equations']>=30 and report['nested_math_paragraphs']==0
    (REVIEW/'08_验证记录/Word结构验证.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(report)
def render():
    import pymupdf
    from PIL import Image,ImageDraw
    pdf=pymupdf.open(REVIEW/'08_验证记录/论文渲染.pdf')
    out=REVIEW/'08_验证记录/论文逐页';out.mkdir(exist_ok=True)
    report=[]
    for i,page in enumerate(pdf):
        page.get_pixmap(matrix=pymupdf.Matrix(1.5,1.5)).save(out/f'page-{i+1}.png')
        text=page.get_text()
        report.append({'page':i+1,'characters':len(text),'start':text[:130],'end':text[-100:]})
    for batch in range((len(pdf)+5)//6):
        im=Image.new('RGB',(1200,1700),'#cccccc');draw=ImageDraw.Draw(im)
        for j in range(6):
            n=batch*6+j+1
            if n>len(pdf):break
            v=Image.open(out/f'page-{n}.png').convert('RGB');v.thumbnail((390,795))
            xx=(j%3)*400;yy=(j//3)*850
            im.paste(v,(xx+(400-v.width)//2,yy+30));draw.text((xx+10,yy+8),f'Page {n}',fill='black')
        im.save(out/f'总览-{batch+1}.jpg',quality=90)
    (out/'页内容检查.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print('PDF_PAGES',len(pdf));print(json.dumps(report,ensure_ascii=False,indent=2))
def originals():
    entries=json.loads((REVIEW/'00_原始文件清单/原始文件SHA256.json').read_text(encoding='utf-8-sig'))
    assert all(hashlib.sha256(Path(r['Path']).read_bytes()).hexdigest().upper()==r['Hash'] for r in entries)
    (REVIEW/'08_验证记录/原始文件未修改.txt').write_text(f'逐一复核{len(entries)}个原始文件，SHA256全部保持不变。\n',encoding='utf-8')
if __name__=='__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    if '--render' in sys.argv:render()
    else:sheets();docx();originals()
