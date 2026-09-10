"""C题微网：数据审计、因果预测、情景日前 LP 和 MPC。单位 kWh / 元。"""
from pathlib import Path
import json, hashlib, shutil, datetime as dt
from functools import lru_cache
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.optimize import linprog
from scipy.sparse import lil_matrix
from scipy.interpolate import PchipInterpolator

ROOT=Path(__file__).resolve().parents[1]
DT=1/6; T=144; SEED=20260910
DATES=pd.date_range('2025-01-01',periods=365)
METHODS=['mean','ma7','ma14','ma30','ewma','ridge']
def dump(obj,path):
    Path(path).write_text(json.dumps(obj,ensure_ascii=False,indent=2,default=lambda x:x.item() if isinstance(x,np.generic) else str(x)),encoding='utf-8')
def time_label(t):
    return f'{t//6:02d}:{t%6*10:02d}'
def interval(t,end=None):return time_label(t)+'-'+time_label(t+1 if end is None else end)
def align_time(label):
    # 原模板最后列统一勘误为本日 00:00--00:10；其他列按左端精确匹配。
    left=str(label).split('-')[0].replace('+1','')
    h,m=map(int,left.split(':')[:2]); return (h*60+m)//10
def _minute(v):
    if isinstance(v,dt.time):return v.hour*60+v.minute
    s=str(v);return 1440 if '+1' in s else int(s.split(':')[0])*60+int(s.split(':')[1])
def load_data():
    if not (ROOT/'input'/'附件1.xlsx').exists():
        raise FileNotFoundError('请将附件1至4及附件5模板放入本复现目录的input文件夹')
    w1=load_workbook(ROOT/'input'/'附件1.xlsx',data_only=True)
    rows=list(w1.active.values); a=np.asarray([r[1:4] for r in rows[1:]],float)
    assert [_minute(r[0]) for r in rows[1:]]==list(range(10,1441,10))
    def matrix(file,sheet=None):
        w=load_workbook(ROOT/'input'/file,data_only=True);s=w[sheet] if sheet else w.active
        r=list(s.values)
        assert [_minute(x) for x in r[0][1:]]==list(range(10,1441,10))
        assert list(pd.to_datetime([x[0] for x in r[1:]]))==list(DATES)
        return np.asarray([x[1:] for x in r[1:]],float)
    L=matrix('附件2.xlsx','小区负载');PV=matrix('附件2.xlsx','光伏发电实际功率');P=matrix('附件4.xlsx')
    w3=load_workbook(ROOT/'input'/'附件3.xlsx',data_only=True)
    F=np.full((365,4,24),np.nan);mapping=[];seen=set();last=None
    for row in list(w3.active.values)[1:]:
        if row[0] not in [None,'']:last=pd.Timestamp(row[0])
        hour=int(str(row[1]).split(':')[0]); day=(last-DATES[0]).days;assert hour in [0,6,12,18]
        assert (day,hour) not in seen;seen.add((day,hour));F[day,hour//6]=row[2:26]
        for k in range(24):mapping.append([str(last+pd.Timedelta(hours=hour)),k+1,str(last+pd.Timedelta(hours=hour+k+1))])
    assert len(seen)==1460
    pd.DataFrame(mapping,columns=['发布时间','预报第几小时','目标时间']).to_csv(ROOT/'audit'/'预报时间映射.csv',index=False,encoding='utf-8-sig')
    stats=[]
    for name,v,unit in [('附件1电价',a[:,0],'元/kWh'),('附件1负载',a[:,1],'kW'),('附件1光伏',a[:,2],'kW'),('附件2负载',L,'kW'),('附件2光伏',PV,'kW'),('附件3预报',F,'kW'),('附件4电价',P,'元/kWh')]:
        stats.append(dict(数据=name,形状=str(v.shape),单位=unit,数量=v.size,缺失=int(np.isnan(v).sum()),非有限=int((~np.isfinite(v)).sum()),负值=int((v<0).sum()),最小=float(v.min()),最大=float(v.max()),均值=float(v.mean())))
        assert np.isfinite(v).all() and (v>=0).all(),name
    pd.DataFrame(stats).to_csv(ROOT/'audit'/'数据审计.csv',index=False,encoding='utf-8-sig')
    manifest={f.name:hashlib.sha256(f.read_bytes()).hexdigest() for f in (ROOT/'input').iterdir()}
    dump(manifest,ROOT/'audit'/'输入SHA256.json')
    tm=[];templates={}
    for f in (ROOT/'input').glob('result*.xlsx'):
        w=load_workbook(f);s=w['计划购电量'];labels=[s.cell(i,1).value for i in range(2,146)] if f.stem=='result1' else [s.cell(1,j).value for j in range(2,146)]
        ids=[align_time(x) for x in labels];assert sorted(ids)==list(range(144))
        for i,(label,t) in enumerate(zip(labels,ids)):tm.append([f.name,i+1,label,interval(t),time_label(t+1),t,'原末列勘误' if t==0 else '按左端匹配'])
        templates[f.name]={s.title:[s.max_row,s.max_column] for s in w}
    pd.DataFrame(tm,columns=['文件','数据序号','原模板标签','规范区间','附件右端采样时刻','内部索引','说明']).to_csv(ROOT/'audit'/'模板逐列时间映射.csv',index=False,encoding='utf-8-sig')
    dump(templates,ROOT/'audit'/'模板结构.json')
    np.savez_compressed(ROOT/'results'/'data.npz',a=a,L=L,PV=PV,P=P,F=F)
    return dict(a=a,L=L,PV=PV,P=P,F=F)

def forecast(y,d,method,fallback):
    """只能访问 y[:d]；不同候选均遵循相同截止点。"""
    if d==0:return fallback.copy()
    h=y[:d]
    if method=='mean':v=h.mean(axis=0)
    elif method.startswith('ma'):v=h[-int(method[2:]):].mean(axis=0)
    elif method=='ewma':
        weight=.8**np.arange(min(d,30)-1,-1,-1);v=np.average(h[-30:],axis=0,weights=weight)
    else:
        if d<12:return forecast(y,d,'ma7',fallback)
        def feat(j):
            return np.stack([np.ones(T),y[j-1]/5000,y[j-7]/5000,y[max(0,j-7):j].mean(axis=0)/5000,
                             np.full(T,np.sin(2*np.pi*j/7)),np.full(T,np.cos(2*np.pi*j/7))],axis=1)
        idx=range(max(7,d-60),d);X=np.concatenate([feat(j) for j in idx]);z=y[list(idx)].reshape(-1)/5000
        reg=np.eye(X.shape[1])*1.;reg[0,0]=1e-6
        coef=np.linalg.solve(X.T@X+reg,X.T@z);v=feat(d)@coef*5000
        # 用历史范围裁剪外推，禁止用当日实际值设定上下界。
        v=np.clip(v,0,np.maximum(h[-30:].max(axis=0)*1.3,1.))
    return np.maximum(v,0)
def build_forecasts(data):
    out={}
    for m in METHODS:
        out[m]={}
        for k,col in [('L',1),('PV',2)]:
            out[m][k]=np.stack([forecast(data[k],d,m,data['a'][:,col]) for d in range(365)])
    return out
def metrics(pred,actual):
    e=pred-actual;mask=actual>max(1,float(actual.max())*.01)
    return dict(MAE=float(np.abs(e).mean()),RMSE=float(np.sqrt((e*e).mean())),nMAE=float(np.abs(e).mean()/max(actual.mean(),1)),MAPE=float(np.mean(np.abs(e[mask]/actual[mask]))*100))

@lru_cache(maxsize=128)
def lp_structure(n,k,eta):
    # G[n], 各情景(C,D,R,W,S_next)[5n], terminal_abs[k], optional plus/minus[2n]
    nvar=n+k*5*n+k+2*n;A=lil_matrix((2*n*k+n,nvar));b=np.zeros(2*n*k+n)
    for s in range(k):
        off=n+s*5*n
        for t in range(n):
            r=2*n*s+t;A[r,t]=1;A[r,off+t]=-1;A[r,off+n+t]=1;A[r,off+2*n+t]=1;A[r,off+3*n+t]=-1
            r=2*n*s+n+t;A[r,off+4*n+t]=1;A[r,off+t]=-eta;A[r,off+n+t]=1/eta
            if t:A[r,off+4*n+t-1]=-1
    plus=n+k*5*n+k
    for t in range(n):A[2*n*k+t,t]=1;A[2*n*k+t,plus+t]=-1;A[2*n*k+t,plus+n+t]=1
    U=lil_matrix((2*k,nvar))
    for s in range(k):
        z=n+k*5*n+s;end=n+s*5*n+5*n-1
        U[2*s,end]=1;U[2*s,z]=-1;U[2*s+1,end]=-1;U[2*s+1,z]=-1
    return A.tocsr(),U.tocsr()

def optimize_day(net_scenarios,p,s0,eta=.9,capacity=12000,emergency=5,terminal=.5,old=None,fee=.5,refund=True,cyclic=False):
    N=np.atleast_2d(net_scenarios);k,n=N.shape;p=np.asarray(p);A,U=lp_structure(n,k,eta)
    nv=A.shape[1];c=np.zeros(nv);b=np.zeros(A.shape[0]);bu=np.tile([capacity*.5,-capacity*.5],k)
    weights=np.ones(k)/k;c[:n]=p if old is None else 0
    bounds=[(0,None)]*nv;eps=1e-7
    for s in range(k):
        off=n+s*5*n;b[2*n*s:2*n*s+n]=N[s];b[2*n*s+n]=s0
        c[off:off+2*n]=eps/k;c[off+2*n:off+3*n]=emergency*p*weights[s]
        c[n+k*5*n+s]=terminal*p.mean()*weights[s]
        for t in range(n):
            bounds[off+t]=(0,5000*DT);bounds[off+n+t]=(0,5000*DT)
            bounds[off+4*n+t]=(capacity*.1,capacity*.9)
            if cyclic:bounds[off+2*n+t]=(0,0)
        if cyclic:bounds[off+5*n-1]=(s0,s0)
    plus=n+k*5*n+k
    if old is None:
        # 删除交易关系行；辅助增减量固定零。
        Ause=A[:2*n*k];buse=b[:2*n*k]
        for t in range(2*n):bounds[plus+t]=(0,0)
    else:
        b[2*n*k:]=old;Ause=A;buse=b
        c[plus:plus+n]=(1+fee)*p;c[plus+n:]=(-1+fee if refund else fee)*p
    sol=linprog(c,A_ub=U,b_ub=bu,A_eq=Ause,b_eq=buse,bounds=bounds,method='highs',options={'dual_feasibility_tolerance':1e-8,'primal_feasibility_tolerance':1e-8})
    if not sol.success:raise RuntimeError(sol.message)
    off=n;C=sol.x[off:off+n];D=sol.x[off+n:off+2*n]
    assert np.max(np.minimum(C,D))<1e-5,'同时充放电'
    return dict(G=np.maximum(sol.x[:n],0),C=C,D=D,S=np.r_[s0,sol.x[off+4*n:off+5*n]],objective=sol.fun)

def recourse(g,net,s,eta=.9,capacity=12000):
    """当前时段测量到达后进行局部能量平衡；绝不读取下一时段。"""
    surplus=g-net
    if surplus>=0:
        C=min(surplus,5000*DT,max(0,(capacity*.9-s)/eta));D=R=0.;W=surplus-C
    else:
        D=min(-surplus,5000*DT,max(0,(s-capacity*.1)*eta));C=W=0.;R=-surplus-D
    return C,D,R,W,s+eta*C-D/eta
def interpolate_hourly_forecast(values,anchor,method='linear'):
    x=np.arange(25);y=np.r_[anchor,values];q=np.arange(1,145)/6
    v=np.interp(q,x,y) if method=='linear' else PchipInterpolator(x,y)(q)
    return np.maximum(v,0)
def scenarios(pred,residuals,margin=.5):
    # 最近历史预测残差的真实整日曲线按全天净能量排序选三条代表；不读取目标日。
    if len(residuals)<3:return pred[None,:]
    r=residuals[-30:];order=np.argsort(r.sum(axis=1));ix=np.rint(np.array([.15,.5,.85])*(len(r)-1)).astype(int)
    return pred[None,:]+r[order[ix]]+margin*r.std(axis=0)[None,:]
def merge_emergency_intervals(r,tol=1e-8):
    out=[];start=None
    for t in range(len(r)+1):
        active=t<len(r) and r[t]>tol
        if active and start is None:start=t
        if not active and start is not None:out.append((interval(start,t),float(r[start:t].sum())));start=None
    assert abs(sum(x[1] for x in out)-r.sum())<1e-6
    return out

def run_policy(data,fl,fp,margin=.5,terminal=.5,eta=.9,capacity=12000,emergency=5,updates=(),q3=False,price_actual=False,price_forecast=False,fee=.5,refund=True,interp='linear',days=None,s0=None,tag='',save=True):
    if days is None:days=range(365)
    days=list(days);s=capacity*.5 if s0 is None else s0;out={k:[] for k in ['G0','G','C','D','R','W','S','plus','minus','PVhat','Lhat','Pdecision']};daily=[];versions=[]
    netactual=(data['L']-data['PV'])*DT
    resid=netactual-(fl-fp)*DT
    for d in days:
        # 正式评价从2月开始。各问题3策略在1月统一按问题2预运行，确保2月初状态可比。
        q3_active=q3 and not (days[0]==0 and d<31)
        p=data['P'][d] if price_actual else data['a'][:,0]
        pp=forecast(data['P'],d,'ma7',data['a'][:,0]) if price_forecast else p
        pv0=fp[d].copy()
        if q3_active:pv0=interpolate_hourly_forecast(data['F'][d,0],data['PV'][d-1,-1] if d else 0,interp)
        pred=(fl[d]-pv0)*DT
        # 问题3用同一发布时刻过去日期的预报残差，避免把日前误差套在临近预报。
        if q3_active:
            hist=[]
            for j in range(max(0,d-30),d):
                pj=interpolate_hourly_forecast(data['F'][j,0],data['PV'][j-1,-1] if j else 0,interp)
                hist.append(netactual[j]-(fl[j]-pj)*DT)
            rs=np.asarray(hist).reshape(-1,T)
        else:rs=resid[:d]
        m=.5 if d<31 and days[0]==0 else margin;term=.5 if d<31 and days[0]==0 else terminal
        N=scenarios(pred,rs,m)
        plan=optimize_day(N,pp,s,eta,capacity,emergency,term);g0=plan['G'].copy();g=g0.copy()
        C=np.zeros(T);D=C.copy();R=C.copy();W=C.copy();plus=C.copy();minus=C.copy();S=np.zeros(T+1);S[0]=s;ph=pv0.copy();lh=fl[d].copy()
        vv=[g0.copy()]
        for t in range(T):
            if q3_active and t in [int(h*6) for h in updates]:
                hour=t//6
                pvnew=interpolate_hourly_forecast(data['F'][d,hour//6],data['PV'][d,t-1],interp)[:T-t]
                rnow=(data['L'][d,max(0,t-6):t]-fl[d,max(0,t-6):t]).mean()
                lnew=np.maximum(0,fl[d,t:]+rnow*np.exp(-np.arange(T-t)/36))
                prednew=(lnew-pvnew)*DT
                rh=[]
                for j in range(max(0,d-30),d):
                    hp=interpolate_hourly_forecast(data['F'][j,hour//6],data['PV'][j,t-1],interp)[:T-t]
                    lr=(data['L'][j,max(0,t-6):t]-fl[j,max(0,t-6):t]).mean()
                    hl=np.maximum(0,fl[j,t:]+lr*np.exp(-np.arange(T-t)/36))
                    rh.append(netactual[j,t:]-(hl-hp)*DT)
                N=scenarios(prednew,np.asarray(rh).reshape(-1,T-t),m)
                pl=optimize_day(N,pp[t:],s,eta,capacity,emergency,term,old=g[t:],fee=fee,refund=refund)
                delta=pl['G']-g[t:];plus[t:]+=np.maximum(delta,0);minus[t:]+=np.maximum(-delta,0);g[t:]=pl['G'];ph[t:]=pvnew;lh[t:]=lnew
                vv.append(g.copy())
            C[t],D[t],R[t],W[t],s=recourse(g[t],netactual[d,t],s,eta,capacity);S[t+1]=s
        base=float(p@g0);inc=float((1+fee)*p@plus);cancel=float(fee*p@minus);ref=float(p@minus) if refund else 0.;em=float(emergency*p@R)
        total=base+inc+cancel-ref+em
        daily.append(dict(day=d,date=str(DATES[d].date()),G0=g0.sum(),G=g.sum(),R=R.sum(),W=W.sum(),base=base,increase=inc,cancel=cancel,refund=ref,adjustment=inc+cancel-ref,emergency=em,total=total,S0=S[0],Send=S[-1],C=C.sum(),D=D.sum(),absolute_adjustment=(plus+minus).sum(),surcharge=fee*float(p@(plus+minus))))
        for key,val in [('G0',g0),('G',g),('C',C),('D',D),('R',R),('W',W),('S',S),('plus',plus),('minus',minus),('PVhat',ph),('Lhat',lh),('Pdecision',pp)]:out[key].append(val)
        versions.append(vv)
        if tag and (d%60==0 or d==days[-1]):print(tag,str(DATES[d].date()),'cost',round(total,2),flush=True)
    out={k:np.asarray(v) for k,v in out.items()};out['days']=np.array(days);out['eta']=np.array(eta);out['capacity']=np.array(capacity);out['multiplier']=np.array(emergency);out['fee']=np.array(fee);out['refund_flag']=np.array(refund)
    df=pd.DataFrame(daily)
    if save and tag:
        np.savez_compressed(ROOT/'results'/f'{tag}.npz',**out);df.to_csv(ROOT/'results'/f'{tag}_daily.csv',index=False,encoding='utf-8-sig')
        dump({str(d):[v.tolist() for v in vv] for d,vv in zip(days,versions) if d in [78,171,265,354]},ROOT/'results'/f'{tag}_versions.json')
    return out,df

def validate_solution(out,data,df,price_actual=False):
    ids=out['days'];p=data['P'][ids] if price_actual else np.broadcast_to(data['a'][:,0],(len(ids),T))
    G,C,D,R,W,S=[out[k] for k in ['G','C','D','R','W','S']];eta=float(out['eta']);cap=float(out['capacity']);fee=float(out['fee'])
    bal=G+data['PV'][ids]*DT+D+R-data['L'][ids]*DT-C-W
    expected=(p*out['G0']).sum(axis=1)+(1+fee)*(p*out['plus']).sum(axis=1)+(fee-int(out['refund_flag']))*(p*out['minus']).sum(axis=1)+float(out['multiplier'])*(p*R).sum(axis=1)
    checks={'能量守恒':np.max(np.abs(bal)),'SOC动态':np.max(np.abs(S[:,1:]-S[:,:-1]-eta*C+D/eta)),
            'SOC边界':max(0,cap*.1-S.min(),S.max()-cap*.9),'充放电上限':max(0,C.max()-5000*DT,D.max()-5000*DT),
            '非负':max(0,-min(v.min() for v in [G,C,D,R,W])), '同时充放电':np.minimum(C,D).max(),
            '跨日连续':np.max(np.abs(S[1:,0]-S[:-1,-1])) if np.all(np.diff(ids)==1) and len(ids)>1 else 0.,
            '费用独立核算':np.max(np.abs(expected-df.total.to_numpy())),
            '有效计划交易恒等式':np.max(np.abs(G-out['G0']-out['plus']+out['minus'])),
            '紧急合并守恒':max(abs(sum(v for _,v in merge_emergency_intervals(r))-r.sum()) for r in R),
            '有限数值':float(not all(np.isfinite(x).all() for x in [G,C,D,R,W,S]))}
    result=[dict(检查=k,误差=float(v),状态='PASS' if v<1e-6 else 'FAIL') for k,v in checks.items()]
    assert all(x['状态']=='PASS' for x in result),result
    return result
