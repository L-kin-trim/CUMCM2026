"""将LaTeX转换为MathML，再通过Microsoft官方XSL转换为Word OMML。"""
from pathlib import Path
import re,sys,copy,json
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'公式依赖'))
from latex2mathml.converter import convert
from lxml import etree
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT

XSL=Path('C:/Program Files/Microsoft Office/root/Office16/MML2OMML.XSL')
if not XSL.exists():raise FileNotFoundError('需要Microsoft Office MML2OMML.XSL以转换原生公式')
TRANSFORM=etree.XSLT(etree.parse(str(XSL)))
FORMULAS=[]
def math(tex,display=False):
    node=TRANSFORM(etree.fromstring(convert(tex,display='block' if display else 'inline').encode())).getroot()
    # Office旧版XSL会把widehat转换为上下限结构，并把overline变成基线横杠。
    # 规范化为真正的重音与上划线，保持LaTeX的数学含义。
    for v in list(node.iter(qn('m:limUpp'))):
        lim=v.find(qn('m:lim'));e=v.find(qn('m:e'))
        if lim is not None and ''.join(lim.itertext()).strip()=='^':
            acc=OxmlElement('m:acc');pr=OxmlElement('m:accPr');ch=OxmlElement('m:chr');ch.set(qn('m:val'),'\u0302');pr.append(ch);acc.append(pr);acc.append(copy.deepcopy(e));v.getparent().replace(v,acc)
    for v in list(node.iter(qn('m:acc'))):
        ch=v.find('.//'+qn('m:chr'))
        if ch is not None and ch.get(qn('m:val')) in ['\u2015','\u00af']:
            bar=OxmlElement('m:bar');pr=OxmlElement('m:barPr');pos=OxmlElement('m:pos');pos.set(qn('m:val'),'top');pr.append(pos);bar.append(pr);bar.append(copy.deepcopy(v.find(qn('m:e'))));v.getparent().replace(v,bar)
    FORMULAS.append(tex)
    return copy.deepcopy(node)
def inline(p,text):
    for i,part in enumerate(re.split(r'\$([^$]+)\$',text)):
        if not part:continue
        if i%2:p._p.append(math(part))
        else:p.add_run(part)
    return p
def paragraph(doc,text,indent=True):
    p=doc.add_paragraph();inline(p,text)
    p.paragraph_format.first_line_indent=Cm(.74) if indent else Cm(0)
    return p
def equation(doc,tex,num):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(5)
    omath=math(tex,True)
    if omath.tag==qn('m:oMathPara'):
        mp=omath
    else:
        mp=OxmlElement('m:oMathPara');mp.append(omath)
    p._p.append(mp);p.add_run(f'  （{num}）')
    return p
def table(doc,headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if widths:
        for c,w in zip(t.columns,widths):c.width=Cm(w)
    for cell,txt in zip(t.rows[0].cells,headers):inline(cell.paragraphs[0],str(txt))
    for row in rows:
        for cell,txt in zip(t.add_row().cells,row):inline(cell.paragraphs[0],str(txt))
    for i,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if i==0:pr.append(OxmlElement('w:tblHeader'))
        for j,c in enumerate(row.cells):
            if widths:c.width=Cm(widths[j])
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcp=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
            tcp.append(borders)
            mar=OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                b=OxmlElement('w:'+edge);b.set(qn('w:w'),'65');b.set(qn('w:type'),'dxa');mar.append(b)
            tcp.append(mar)
            if i==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EAEAEA');tcp.append(sh)
            for p in c.paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.05
                for r in p.runs:r.font.size=Pt(9);r.bold=i==0
    return t
def setup():
    d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Cm(2.55)
    s.header_distance=s.footer_distance=Cm(1.25)
    for name,size,font in [('Normal',10.5,'宋体'),('Title',18,'黑体'),('Heading 1',14,'黑体'),('Heading 2',12,'黑体'),('Heading 3',10.5,'黑体')]:
        st=d.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),font)
        st.paragraph_format.space_after=Pt(5);st.paragraph_format.line_spacing=1.15
    d.styles['Normal'].paragraph_format.widow_control=True
    for name in ['Heading 1','Heading 2','Heading 3']:d.styles[name].paragraph_format.keep_with_next=True
    footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    for elem in list(d.styles.element.iter(qn('w:pBdr'))):elem.getparent().remove(elem)
    d.core_properties.author='';d.core_properties.last_modified_by='';d.core_properties.title='微网与外部电网电力调控策略的预测优化模型'
    return d

if __name__=='__main__':
    d=setup();d.add_heading('公式转换验证',0)
    equation(d,r'S_{t+1}=S_t+\eta_c C_t-\frac{D_t}{\eta_d}',1)
    equation(d,r'\min\sum_{t=0}^{143}p_tG_t+\frac{1}{3}\sum_{s=1}^{3}\sum_{t=0}^{143}5p_tR_t^{(s)}',2)
    paragraph(d,r'日前计划 $G_t^0$ 满足 $0\leq C_t,D_t\leq 5000/6$，且 $\widehat{L}_{d,t}$ 只依赖过去数据。')
    d.save(ROOT.parent/'08_验证记录/公式测试.docx')
