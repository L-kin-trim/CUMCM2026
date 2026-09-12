"""按用户提供的范例排版；数值读取上一轮已经验证的结果，不改写既有成果。"""
from pathlib import Path
import sys, json, csv, copy, re
HERE=Path(__file__).resolve().parent
WORK=HERE.parent.parent
SRC=WORK/'审查与完善/07_复现代码与数据/src'
sys.path.insert(0,str(SRC))
from word_math import *
from docx.enum.text import WD_BREAK, WD_LINE_SPACING
from docx.enum.section import WD_SECTION_START
R=ROOT/'results'
def rows(name):
    with (R/name).open(encoding='utf-8-sig') as f:return list(csv.DictReader(f))
summary={r['model']:r for r in rows('年度汇总.csv')}
def val(tag,key):return float(summary[tag][key])
def f(x,n=2):return f'{float(x):,.{n}f}'
q1=json.loads((R/'q1_summary.json').read_text(encoding='utf-8'))
cert=json.loads((ROOT/'计算结果/summary.json').read_text(encoding='utf-8'))
typ=json.loads((R/'指定日期完整结果.json').read_text(encoding='utf-8'))
doc=setup();notes=[];n_eq=0;n_tab=0;n_fig=0
for s in doc.sections:s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Cm(2.5)
for name,size,cn,bold in [('Normal',12,'宋体',False),('Title',18,'黑体',False),('Heading 1',16,'黑体',False),('Heading 2',14,'黑体',False),('Heading 3',12,'宋体',True)]:
    st=doc.styles[name];st.font.size=Pt(size);st.font.bold=bold
    st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),cn)
    pf=st.paragraph_format;pf.space_before=Pt(0);pf.space_after=Pt(0);pf.line_spacing=Pt(23.4)
    if name!='Normal':pf.first_line_indent=Pt(0)
    if name=='Heading 1':pf.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf.space_before=Pt(12);pf.space_after=Pt(8)
    if name=='Heading 2':pf.space_before=Pt(6);pf.space_after=Pt(3)
    if name=='Heading 3':pf.space_before=Pt(4)
    if name=='Title':pf.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf.space_before=Pt(7);pf.space_after=Pt(12);pf.line_spacing=Pt(31.2)
doc.styles['Normal'].paragraph_format.first_line_indent=Pt(24)
doc.styles['Normal'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
doc.settings.element.append(OxmlElement('w:doNotAutoHyphenate')) if False else None
def p(text,indent=True):
    notes.append(text);pp=paragraph(doc,text,False);pp.paragraph_format.first_line_indent=Pt(24 if indent else 0);return pp
def h(text,level=1):notes.append('#'*(level+1)+' '+text);return doc.add_heading(text,level)
def br():doc.add_page_break()
def eq(tex):
    global n_eq
    n_eq+=1;notes.append('$$\n'+tex+'\n$$')
    pp=equation(doc,tex,n_eq);pp.paragraph_format.first_line_indent=Pt(0)
    pp.paragraph_format.line_spacing=1.0;pp.paragraph_format.space_before=Pt(5);pp.paragraph_format.space_after=Pt(5)
    # 不采用固定行高，避免分式、上下标和矩阵被裁切。
    return pp
def caption(text,before=True):
    pp=p(text,False);pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    pp.paragraph_format.line_spacing=Pt(18);pp.paragraph_format.space_before=Pt(5);pp.paragraph_format.space_after=Pt(3);pp.paragraph_format.keep_with_next=before
    for rr in pp.runs:rr.font.size=Pt(10.5)
    return pp
def tab(title,heads,data,widths=None,size=10.5):
    global n_tab
    n_tab+=1;caption(f'表5-{n_tab} {title}')
    notes.append('| '+' | '.join(heads)+' |\n| '+' | '.join(['---']*len(heads))+' |\n'+'\n'.join('| '+' | '.join(map(str,row))+' |' for row in data))
    t=table(doc,heads,data,widths or [16/len(heads)]*len(heads))
    for i,row in enumerate(t.rows):
        for c in row.cells:
            tcp=c._tc.get_or_add_tcPr()
            for child in list(tcp):
                if child.tag in [qn('w:shd'),qn('w:tcBorders')]:tcp.remove(child)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','bottom','left','right']:
                b=OxmlElement('w:'+edge);on=(edge=='top' and i==0) or (edge=='bottom' and (i==0 or i==len(t.rows)-1))
                b.set(qn('w:val'),'single' if on else 'nil');b.set(qn('w:sz'),'12' if (i==len(t.rows)-1 or edge=='top') else '6');b.set(qn('w:color'),'000000');borders.append(b)
            tcp.append(borders)
            for pp in c.paragraphs:
                pp.paragraph_format.first_line_indent=Pt(0);pp.paragraph_format.line_spacing=Pt(16);pp.paragraph_format.keep_with_next=i<len(t.rows)-1
                for rr in pp.runs:rr.font.size=Pt(size)
    return t
def fig(path,title,width=15):
    global n_fig
    n_fig+=1;pp=doc.add_paragraph();pp.alignment=WD_ALIGN_PARAGRAPH.CENTER;pp.paragraph_format.first_line_indent=Pt(0);pp.paragraph_format.line_spacing=1;pp.paragraph_format.keep_with_next=True
    pp.add_run().add_picture(str(path),width=Cm(width));caption(f'图5-{n_fig} {title}',False)
def step(label,text):
    pp=p(label);pp.runs[0].bold=True;pp.paragraph_format.keep_with_next=True;p(text)

doc.add_heading('基于滚动预测与储能协同调度的\n微网购电策略优化研究',0)
ab=p('摘　要',False);ab.alignment=WD_ALIGN_PARAGRAPH.CENTER;ab.paragraph_format.space_after=Pt(12)
for rr in ab.runs:rr.font.size=Pt(16);rr.font.name='Times New Roman';rr._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'黑体')
p('本文针对微网与外部电网的电力调控问题，建立能量平衡与储能状态约束下的购电优化模型，研究确定性日计划、负载与光伏变化、预报更新和波动电价四种条件下的调度策略。通过线性规划、滚动预测和历史残差情景，协调常规购电、储能充放电与紧急补购。')
p(f'问题一中，建立以日购电费用最小为目标的线性规划模型，在首尾储电量相等的条件下，求得日购电量{f(cert["purchase"])} kWh，最低费用{f(cert["cost"])}元，较无储能方案节省{f(cert["savings_pct"])}%。通过原始与对偶目标的一致性检验验证该结果。')
p(f'问题二中，采用历史数据训练岭回归预测器，以三条残差情景制定日前计划，并根据当前净负荷执行储能调节。2月至12月总费用为{f(val("q2","total")/10000)}万元。问题三中，建立考虑增购、取消和紧急购电的滚动调整模型；采用0、6、12、18时全部预报时，总费用为{f(val("q3_D","total")/10000)}万元，较仅用0时预报降低4.61%，紧急购电量降低87.18%。')
p(f'问题四中，将分时固定电价推广为波动电价。在当天价格曲线已知的条件下，日前计划与滚动调整方案的费用分别为{f(val("q42","total")/10000)}万元和{f(val("q43_D","total")/10000)}万元；仅用历史价格预测时，对应费用为{f(val("q42_causal_price","total")/10000)}万元和{f(val("q43_causal_price","total")/10000)}万元。两类结果分别反映不同的信息条件。')
p('结果表明，储能移峰与预报更新在所测策略中能够降低购电支出。模型通过能量守恒、容量、功率和费用核算检验；除问题一外，结果为给定策略的历史回放费用，其适用范围受预测误差、结算约定与实时控制规则限制。')
kw=p('关键词：');kw.runs[0].bold=True;kw.add_run('微网；储能调度；线性规划；滚动预测；情景优化')
br()
h('第1章　问题重述')
h('1.1 问题背景',2)
p('微网由小区负载、光伏发电设备和储能设备组成，并通过外部电网补充电能。光伏出力与用电需求在时间上的差异，使得单纯依靠即时购电难以充分利用低价电能和本地发电。储能设备能够在低价或光伏富余时充电，在高价或供能不足时放电，从而改变外网购电的时间分布。')
p('本研究考虑的调控目标是在持续满足小区负载的前提下，尽可能降低购电费用。系统每10分钟记录一次功率，储能最大容量为12000 kWh，允许储电量为1200至10800 kWh，最大充放电功率为5000 kW，初始储电量为6000 kWh。预测误差与电价变化进一步影响日前计划和后续调整，需要在统一的能量约束下进行分析。')
h('1.2 问题要求',2)
for text in [
'问题1：每天电价和负载相同，已知一天的光伏预测，在0时制定日购电计划，并满足日初与日末储电量相等。依据附件1给出指定购电时段、六个储能统计区间及全天费用。',
'问题2：负载和光伏功率随时间变化，每天0时制定计划，其余常规费用按计划结算。供能不足部分按当时电价的5倍紧急购电，需利用附件2给出全年计划及指定日期结果。',
'问题3：每天0、6、12、18时获得未来24小时的整点光伏预报，允许调整未执行计划，并计入增购或取消交易的相关费用。利用附件3建立滚动调整模型，分析其他预报时刻的应用价值。',
'问题4：利用附件4的波动电价重新计算问题2和问题3，分别给出日前计划和滚动调整结果，并说明电价信息在决策时是否可用。']:
    p(text)

h('第2章　问题分析')
h('2.1 问题1的分析',2)
p('问题1的核心在于将电能的跨时段转移转化为带储能状态的费用最小化问题。由于负载、电价和光伏预测均已给定，购电量与充放电量可以作为连续决策变量。功率约束限制单个时段可转移的电量，容量约束限制跨时段积累的电量，首尾状态相等则防止通过消耗初始库存获得虚假的低费用。')
p('在变量定义一致的条件下，供需平衡和储能动态均可写为线性关系。因此，采用线性规划直接求解，并以独立模型的对偶目标验证最优性。同时结合日内曲线解释低价充电、高价放电及部分时段仍需购电的原因。')
h('2.2 问题2的分析',2)
p('问题2的困难在于日前决策不能使用尚未发生的实际负载和光伏数据。计划过低会增加高价紧急购电，计划过高又可能造成已付费电量弃置；储能则在二者之间提供缓冲。因此，预测误差应进入购电决策，而不宜仅把一条点预测当作确定性输入。')
p('本文先比较不同历史预测方法，再利用完整日残差构造净负荷情景，以共同的日前购电量连接各情景。运行时根据当前观测执行充放电，并将日末状态传递至下一日。这样既能保持决策的时间顺序，又能通过实际账单评价预测与调度的综合效果。')
h('2.3 问题3的分析',2)
p('问题3在日前计划基础上增加了预报更新和交易调整。模型需要解决两个衔接问题：一是将整点预报映射为10分钟序列，二是用更新时刻的实际储电量连接剩余时段的优化。已执行的交易不可回改，尚未执行的计划则需权衡预报改善带来的收益与调整费用。')
p('由于一次调整会影响后续储能状态，各发布时间的作用不能简单相加。本文采用逐步加入6、12、18时更新的嵌套策略进行比较，分析其总费用和紧急购电变化。对于题面未提供预报的其他时刻，只讨论引入条件，不虚构其预测效果。')
h('2.4 问题4的分析',2)
p('波动电价使购电时间的选择同时依赖净负荷和价格判断。附件4给出的是实际发生的价格，不能仅因数据可见就认定当日零点已知全天曲线。因此，应将已知当天价格的条件基准与仅利用历史价格的预测方案分开求解，并统一按实际交易价格结算。')
p('本文沿用问题2和问题3的物理约束与执行规则，只改变决策使用的价格输入。通过并列比较两种信息条件，区分价格预测误差与预报更新的影响，使结果解释与模型假设一致。')

h('第3章　符号说明')
# 符号表单独使用章号，与范例的无编号符号表相同。
t=tab('主要符号及单位',['符号','含义','单位'],[
[r'$d,t$','日期与日内时段索引','—'],[r'$\Delta t$','时段长度，取1/6小时','h'],[r'$L_{d,t},P^{pv}_{d,t}$','小区负载与光伏发电功率','kW'],[r'$E^L_{d,t},E^{pv}_{d,t}$','负载用电量与光伏发电量','kWh'],[r'$p_{d,t},\widehat p_{d,t}$','实际电价与预测电价','元/kWh'],[r'$G^0_{d,t},G^k_{d,t}$','日前计划量与第k次调整后计划量','kWh'],[r'$C_{d,t},D_{d,t}$','母线侧充电量与放电量','kWh'],[r'$S_{d,t}$','时段起点的储电量','kWh'],[r'$R_{d,t},W_{d,t}$','紧急购电量与弃置电量','kWh'],[r'$\eta_c,\eta_d$','充电效率与放电效率','—'],[r'$N_{d,t},N^{(s)}_{d,t}$','实际净负荷电量与情景净负荷电量','kWh'],[r'$\delta^{k,+}_{d,t},\delta^{k,-}_{d,t}$','第k次调整的增购量与减购量','kWh'],[r'$J,\kappa$','购电总费用与终端偏差系数','元；元/kWh']],[4.2,8.8,3])
# 将符号表标题改为不带模型章节编号。
for pp in doc.paragraphs:
    if pp.text=='表5-1 主要符号及单位':pp.runs[0].text='主要符号及单位'
n_tab=0

h('第4章　模型假设')
for text in [
'1. 时间离散假设：每个功率采样值代表此前10分钟区间的平均功率，时间标签为右端点。本文据此转换功率与电量，并在全文保持统一映射。',
'2. 储能效率假设：充电和放电效率分别取90%，往返效率为81%。充电量和放电量均在母线侧计量；若90%指往返效率，则另作敏感性分析。',
'3. 设备与交易假设：忽略储能自放电、老化和额外运维费用，不设置售电收益及外网购电功率上限。允许弃光和已付费富余电量弃置。',
'4. 状态连续假设：问题1首尾储电量均为6000 kWh；问题2至问题4从1月1日的6000 kWh连续运行，不对每天施加首尾相等条件。',
'5. 信息可用性假设：未来实际负载和光伏数据仅用于事后评价。预测使用此前数据，实时执行使用当前观测；问题4分别研究已知当日电价和仅用历史价格的条件。',
'6. 调整结算假设：取消电量退回原款，并支付50%的违约费用；每次增减按上一有效计划计算。题面未明确的退款和多次调整基准通过该约定补充，并检验不退款情形。',
'7. 初始化假设：1月1日无历史样本时，以附件1负载、光伏与电价曲线冷启动；1月用于预运行和参数选择，正式统计区间为2月1日至12月31日。']:
    p(text,False)

h('第5章　模型建立与求解')
h('5.1 数据处理与统一物理约束',2)
h('5.1.1 数据整理与时间映射',3)
p('附件1提供日内电价、负载和光伏预测；附件2提供全年实际负载与光伏；附件3提供分时发布的未来整点预报；附件4提供全年实际电价。各序列统一按日期和日内时段建立索引。数值区域共193152个数据，未发现缺失、负值或非有限数；附件3组内空白日期采用该组首行日期，不作为预报缺失处理。')
eq(r'\Delta t=\frac{1}{6},\qquad E^L_{d,t}=L_{d,t}\Delta t,\qquad E^{pv}_{d,t}=P^{pv}_{d,t}\Delta t')
p('例如10:10采样行对应10:00—10:10区间。结果模板末行或末列的次日标记存在歧义，本文按程序实际映射将其标为本日00:00—00:10，并保留其他列顺序。这一时间解释属于建模约定，不视为官方勘误。四小时储能统计由连续24个10分钟时段求和得到。')
h('5.1.2 能量平衡模型',3)
p('微网母线的供给包括常规购电、光伏发电、储能放电和紧急购电；其去向为负载、储能充电与弃置。引入非负弃置变量后，题目中供电不低于负载的要求可以统一表达为：')
eq(r'G_{d,t}+E^{pv}_{d,t}+D_{d,t}+R_{d,t}=E^L_{d,t}+C_{d,t}+W_{d,t}')
p('弃置电量不产生收入，已经购入的富余电量仍须支付费用。因此，不能用负购电量代替弃置，也不能将全部富余电量当作可售电量。')
fig(HERE/'新增图/01_微网能量流.png','微网能量流及储能双向调节',15.5)
p('能量流图将母线的供给与去向对应到平衡方程。储能同时具有吸收和释放电能两种运行方向，实际时段按状态选择其一；富余弃置不产生售电收入。')
h('5.1.3 储能状态与运行边界',3)
p('充电时，输入电量的一部分转化为储能库存；放电时，为向母线提供规定电量，库存减少量需要考虑效率损失。状态转移与运行边界分别为：')
eq(r'S_{d,t+1}=S_{d,t}+\eta_c C_{d,t}-\frac{D_{d,t}}{\eta_d}')
eq(r'1200\leq S_{d,t}\leq10800,\qquad 0\leq C_{d,t},D_{d,t}\leq\frac{5000}{6}')
eq(r'G_{d,t},R_{d,t},W_{d,t}\geq0,\qquad S_{d+1,0}=S_{d,144}')
p('后续各问均在上述约束下求解。功率上限必须逐时段检查，不能仅凭四小时汇总电量推断设备是否越限；连续运行还需检查相邻日期的首尾状态是否一致。')

h('5.2 问题1的模型建立与求解',2)
p('问题一要求在已知日内曲线的条件下确定购电与储能安排。本文先建立日购电费用目标，再施加首尾储能一致约束，通过线性规划得到完整的10分钟调度，最后按题设格式汇总。')
h('5.2.1 目标函数与边界条件',3)
eq(r'\min J_1=\sum_{t=0}^{143}p_tG_t,\qquad R_t=0,\qquad S_0=S_{144}=6000')
p('决策变量由144个购电量、144个充电量、144个放电量、144个弃置量和145个储电状态组成。结合统一物理约束，模型为线性规划，可采用HiGHS求解器计算[1,2]。为稳定选择解，调度程序加入极小的吞吐惩罚；独立最优性检验则直接最小化真实购电费。')
eq(r'J_{\varepsilon}=J_1+10^{-7}\sum_{t=0}^{143}(C_t+D_t)')
p(r'在正电价且允许弃置的条件下，若某时段同时充放电，将充电减少$a$、放电减少$\eta_c\eta_d a$，储能变化保持不变，母线富余增加$(1-\eta_c\eta_d)a$。通过减少购电或增加弃置可恢复平衡，真实费用不增加而吞吐惩罚下降。因此，存在不同时充放电的最优解，无需在本模型中额外引入整数变量。')
h('5.2.2 求解步骤',3)
step('Step1：数据与参数初始化','将附件1功率换算为时段电量，设置效率、容量、功率及6000 kWh首尾状态，按统一时间映射排列144个区间。')
step('Step2：约束矩阵构造与求解','逐时段建立能量平衡和储能转移等式，设置各变量上下界，求解线性规划并提取购电、充放电、弃置与储电轨迹。')
step('Step3：独立验证与结果汇总','另建仅含真实费用目标的线性规划，比较原始与对偶目标，并用双单纯形和内点方法交叉检查；通过后汇总指定区间和全天指标。')
h('5.2.3 求解结果及分析',3)
tab('问题1的优化结果',['计算项目','数值结果'],[['全天购电量 / kWh',f(cert['purchase'],4)],['全天购电费 / 元',f(cert['cost'],6)],['无储能费用 / 元',f(cert['baseline_cost'],6)],['费用降低比例',f(cert['savings_pct'])+'%'],['原始与对偶目标差 / 元',r'$7.28\times10^{-12}$'],['全天弃置电量 / kWh',f(cert['spill'],4)]],[9,7])
qs=list(q1['six_slots'].items())
tab('问题1指定时段购电量及全天结果',['时间段','购电量','时间段','购电量','时间段','购电量'],[[qs[j][0],f(qs[j][1],4),qs[j+1][0],f(qs[j+1][1],4),qs[j+2][0],f(qs[j+2][1],4)] for j in [0,3]]+[['全天购电量',f(cert['purchase'],4),'全天购电费',f(cert['cost']),'','']],[2.8,2.53,2.8,2.53,2.8,2.54],10)
b=cert['blocks']
tab('问题1储能充放电量与首尾储电量',['时间段','充电量','放电量','时间段','充电量','放电量'],[[b[j]['period'],f(b[j]['charge'],4),f(b[j]['discharge'],4),b[j+1]['period'],f(b[j+1]['charge'],4),f(b[j+1]['discharge'],4)] for j in [0,2,4]]+[['0:00储电量','6000.0000','','24:00储电量','6000.0000','']],[2.8,2.6,2.6,2.8,2.6,2.6],10)
p(f'表中电量单位为kWh、费用单位为元。全天充电量为{f(cert["charge"],4)} kWh，放电量为{f(cert["discharge"],4)} kWh，首尾储电量均为6000 kWh。最大充电功率达到5000 kW，最大放电功率为{f(cert["checks"]["maximum_discharge_kw"],4)} kW，均满足约束。')
fig(ROOT/'计算结果/调度曲线.png','问题1购电与储能调度结果',15)
p('从计算结果可以看出，储能通过低价充电和高价放电降低了外网购电支出。中午即使光伏接近或超过负载，较低电价仍可能使购电充电具有经济性，因此12:00—12:10的购电量不必为零。晚间高价时段的放电受功率和库存限制，部分需求仍须由外网补充。')
p('在首尾储能相等时，全天供给与需求之差应等于充放电损耗，其检验关系为：')
eq(r'\sum_t(G_t+E^{pv}_t-E^L_t-W_t)=(1-\eta_c)\sum_tC_t+(\eta_d^{-1}-1)\sum_tD_t')
p(r'独立计算的全天能量残差小于$10^{-9}$ kWh，两种求解算法给出相同费用。该检验说明充电量大于放电量源于效率损失，而非能量平衡错误。')

h('5.3 问题2的模型建立与求解',2)
h('5.3.1 基于历史数据的功率预测',3)
p('未来负载与光伏未知时，预测必须先于调度。本文比较历史同时间均值、7日、14日、30日移动平均、指数加权移动平均和岭回归六种方法，并采用滚动起点验证[3]。岭回归输入包括前1日、前7日、近7日均值、周周期正余弦和截距，拟合窗口最多60天，功率按5000 kW缩放。')
eq(r'\widehat{\boldsymbol\beta}=(X^{\mathsf T}X+\Lambda)^{-1}X^{\mathsf T}y')
eq(r'\widehat y_{d,t}=5000\,x_{d,t}^{\mathsf T}\widehat{\boldsymbol\beta}')
p('非截距正则系数取1，截距系数取百万分之一；训练不足12天时使用7日均值。预测截断为非负，其上界根据历史范围确定。1月17日至31日用于经济选型，比较6种预测方法、3种裕度和3种终端系数，共54组，选出岭回归、裕度0和终端系数0.5。1月实际预运行仍采用预先固定的参数，避免用月底选型结果改写月初状态。')
pred={(r['method'],r['variable']):r for r in rows('预测精度.csv') if r['period']=='2月至12月外样本'}
tab('不同预测方法的外样本误差',['预测方法','负载MAE','负载RMSE','光伏MAE','光伏RMSE'],[[label]+[f(pred[(m,k)][metric]) for k in ['L','PV'] for metric in ['MAE','RMSE']] for m,label in [('mean','历史均值'),('ma7','7日均值'),('ma14','14日均值'),('ma30','30日均值'),('ewma','指数加权'),('ridge','岭回归')]],[3.4,3.15,3.15,3.15,3.15])
p('表中单位为kW，评价区间为2月至12月。预测误差用于解释运行效果，未用该外样本结果反向选择参数；经济选型以1月验证期间的库存校正费用为依据，因而不要求误差最小的方法必然具有最低购电费用。')
fig(HERE/'新增图/02_典型日预测.png','2025年6月21日的负载与光伏预测对比',15.5)
p('典型日曲线显示，预测能够刻画主要日内变化，但局部偏差仍然存在。负载与光伏偏差共同改变净负荷，因此购电优化还需考虑历史残差情景；单日图形仅作说明，整体精度以外样本误差表为准。')
h('5.3.2 历史残差情景与日前目标函数',3)
p('为描述净负荷的不确定性，保留最近30个完整日的预测残差，按全天残差能量排序，取约15%、50%和85%位置的三条整日曲线。相较于逐时独立扰动，这种构造保留了部分日内相关性，但三个等权情景并不代表已校准的真实概率分布。')
eq(r'\widehat N_{d,t}=(\widehat L_{d,t}-\widehat P^{pv}_{d,t})\Delta t')
eq(r'N^{(s)}_{d,t}=\widehat N_{d,t}+e^{(s)}_{d,t}+m\sigma_{d,t}')
p('日前购电量在三个情景间共享，储能、紧急购电和弃置变量可随情景变化。目标函数同时考虑计划费用、情景紧急费用及终端储能偏差：')
eq(r'\min\left\{\sum_t p_tG_t+\frac{1}{3}\sum_{s=1}^{3}\left[\sum_t5p_tR_t^{(s)}+\kappa z_s\right]\right\}')
eq(r'z_s\geq S_{144}^{(s)}-6000,\quad z_s\geq6000-S_{144}^{(s)},\quad\kappa=0.5\overline p')
p('各情景均满足统一物理约束。这里允许情景内动作依赖整条情景，因此所得规划是用于选择日前购电的两阶段近似；它没有建立完整场景树的非预见性约束，不能视为严格的多阶段最优控制。实际电池执行另按当前观测决定。')
h('5.3.3 实时储能控制与日际衔接',3)
p(r'记当前实际净负荷电量为$N_t$，常规购电抵消净负荷后的余量为$x_t$。富余时优先充电，缺口时优先放电，超过储能调节能力的缺口进行紧急购电：')
eq(r'x_t=G_t-N_t')
eq(r'C_t=\min\{x_t,5000/6,(10800-S_t)/\eta_c\},\quad x_t\geq0')
eq(r'D_t=\min\{-x_t,5000/6,\eta_d(S_t-1200)\},\quad x_t<0')
eq(r'R_t=-x_t-D_t,\quad x_t<0')
p('富余分支中的放电量和紧急购电量为零，剩余富余电量弃置；缺口分支中的充电量和弃置量为零。该规则只使用当前状态与观测，能够满足物理约束，但可能提前消耗后续高价时段需要的储能，因此属于可执行的贪心策略。日末状态连续传递到下一日，不人为重置为6000 kWh。')
h('5.3.4 求解结果及分析',3)
tab('问题2的334天费用与电量',['计算项目','数值结果'],[['计划购电量 / kWh',f(val('q2','G0'))],['计划购电费 / 元',f(val('q2','base'))],['紧急购电量 / kWh',f(val('q2','R'))],['紧急购电费 / 元',f(val('q2','emergency'))],['购电总费用 / 元',f(val('q2','total'))]],[9,7])
p('从费用构成可以看出，紧急购电虽然只占部分电量，却以5倍价格结算，对总费用有显著影响。扩大计划裕度可能减少缺口，但也会增加已付费富余，因此不能仅以紧急购电量最小作为计划目标。指定日期的购电、储能与紧急区间见5.6节，均来自实际连续回放轨迹。')

fig(HERE/'新增图/03_月度紧急购电.png','问题2与问题3策略D的月度紧急购电量',15.5)
p('月度统计将全年紧急购电分解到各月，用于观察风险在时间上的分布及预报更新后的变化。两组策略均按同一实际序列连续运行，图中差异反映预测输入、计划调整和储能状态传递的综合作用。')
h('5.4 问题3的模型建立与求解',2)
h('5.4.1 分时预报更新与剩余时域',3)
p('附件3每次发布24个未来整点预报。本文以发布时间已经观测到的光伏功率作为左端锚点，在线性插值后提取当天未执行时段，并使用相同发布时间的历史预报误差构造残差情景。该处理区分了日前误差与临近预报误差。')
p('负载没有同步发布的外部预报，因此用最近1小时已实现的残差均值修正日前预测，并按6小时尺度衰减。设发布时间为日内第若干小时，修正后的负载预测为：')
eq(r'\widehat L_t^{\tau}=\max\{0,\widehat L_t^0+\overline e_{\tau}\exp[-(t-6\tau)/36]\}')
p(r'其中$\tau$以小时计、$t$为10分钟时段索引。更新时以实际储电量作为剩余模型的初始状态，冻结所有已执行交易，只优化后续计划。本文借鉴滚动时域思想[4]，实施6小时预报触发的购电调整，电池每10分钟继续按实时规则执行。')
fig(HERE/'新增图/04_滚动调整热力图.png','2025年6月21日各次预报触发的计划增减',15.5)
p('热力图中的红色表示相对上一有效计划增购，蓝色表示减购，灰色为更新时刻以前已经执行的区间。图中各次调整仅作用于未来时段，直观体现了已执行交易冻结的约束。')
h('5.4.2 调整费用模型',3)
p('每次调整以此前有效计划为基准，将差额分为增购量和减购量，避免重复计费。按模型假设，增购按1.5倍电价结算；取消部分退回原款并收取50%违约费，因此其净退款为原价的50%。')
eq(r'\delta_t^{k,+}=\max\{G_t^k-G_t^{k-1},0\},\quad\delta_t^{k,-}=\max\{G_t^{k-1}-G_t^k,0\}')
eq(r'J_3=\sum_t p_tG_t^0+\sum_{k,t}(1.5p_t\delta_t^{k,+}-0.5p_t\delta_t^{k,-})+\sum_t5p_tR_t')
eq(r'G_t^{\mathrm{final}}=G_t^0+\sum_k\delta_t^{k,+}-\sum_k\delta_t^{k,-}')
p('将最终常规购电量代入后，总费用还可写成最终常规电费、纯调整附加费与紧急购电费之和：')
eq(r'J_3=\sum_t p_tG_t^{\mathrm{final}}+0.5\sum_{k,t}p_t(\delta_t^{k,+}+\delta_t^{k,-})+\sum_t5p_tR_t')
p('两种写法应给出相同费用。净调整费可能为负，因为其中包含取消交易的退款；纯调整附加费则非负。费用核算时必须保留这一区别。')
h('5.4.3 求解步骤与更新时点比较',3)
step('Step1：制定零点计划','依据0时光伏预报和历史负载预测构造情景，求解当天购电计划，并执行到下一预报时刻。')
step('Step2：更新状态与剩余计划','收到新预报后，读取已实现的储电状态和负载残差；对尚未执行区间重新求解，记录相对上一有效计划的增购、减购及相应费用。')
step('Step3：完成日内执行与年度评价','持续进行10分钟储能调节和紧急补购，将日末状态传递到次日。分别回放只用0时、加入6时、再加入12时及18时的四组策略。')
tab('不同预报更新策略的比较',['策略','预报时刻','总费用 / 元','紧急量 / kWh','净调整费 / 元'],[[z,hours,f(val(k,'total')),f(val(k,'R')),f(val(k,'adjustment'))] for z,hours,k in [('A','0','q3_A'),('B','0、6','q3_B'),('C','0、6、12','q3_C'),('D','0、6、12、18','q3_D')]],[1,3,4.2,3.7,4.1],10)
fig(ROOT/'figures/图12_预报更新策略成本.png','预报更新策略的费用比较',13.8)
p('由A到B、由B到C、由C到D，总费用依次减少377753.91元、182314.59元和114277.05元。采用全部预报的策略D相对A节省4.61%，紧急购电量降低87.18%。因此，在现有数据与结算约定下，保留6、12、18时更新具有经济依据。')
p('上述比较存在状态传递和策略联动，不能解释为每个时点独立的信息价值。对于题面以外的新增时刻，还需取得该时刻实际可用的预报，比较新增调整费、预测改善及储能机会成本。现有结果不能证明任意提高更新频率都会降低费用。')

h('5.5 问题4的模型建立与求解',2)
h('5.5.1 波动价格的信息条件',3)
p('问题4保留前三问的能量和交易约束，将固定分时价格替换为附件4的价格序列。为避免未来价格泄漏，本文设置两类方案：一类假定当天曲线已知，作为与既有结果表对应的条件基准；另一类以此前7天同时间价格均值预测当天价格，作为历史信息方案。')
eq(r'\widehat p_{d,t}=\frac{1}{\min(7,d)}\sum_{j=\max(0,d-7)}^{d-1}p_{j,t},\quad d\geq1')
p('第一天采用附件1电价初始化。历史信息方案在各次日内更新中继续使用0时生成的价格预测，不读取后续实际价格；计划优化使用预测价，账单统一按交易时刻的实际价计算。其他物理边界和储能执行规则保持一致，以便比较。')
h('5.5.2 求解结果及分析',3)
tab('波动电价下的购电结果',['策略','价格信息','总费用 / 元','紧急量 / kWh'],[[label,info,f(val(tag,'total')),f(val(tag,'R'))] for label,info,tag in [('4-2','当日曲线已知','q42'),('4-3','当日曲线已知','q43_D'),('4-2扩展','历史价格预测','q42_causal_price'),('4-3扩展','历史价格预测','q43_causal_price')]],[2.5,4.2,5,4.3])
p('当日曲线已知时，滚动调整方案相对日前方案降低费用4.20%；仅利用历史价格预测时，对应降幅为3.83%。两类条件下，利用额外光伏预报均在本次回放中降低了费用。')
p('历史价格预测使4-2和4-3费用分别比条件基准高0.46%和0.85%。这一差额是给定预测与贪心执行规则下的经验结果，不是严格的信息价值上界。已知价格方案不能直接作为未知未来价格条件下的可实施最优策略。')

fig(HERE/'新增图/05_电价购电储能联动.png','2025年6月21日波动电价下的购电与储能联动',15.5)
p('图示对应当天价格曲线已知条件下的问题4-3策略D。共同时间轴展示价格变化、时段购电量与储电状态的联系；储电量虚线为上下边界。调度还受光伏、负载与调整费影响，不能仅凭价格高低判断某时段是否应购电。')
h('5.6 指定日期的调度结果',2)
p('本节按题设表1、表2与表3的结构给出3月20日、6月21日、9月23日和12月21日的结果。各表电量单位为kWh，费用单位为元；日前计划量与全天计划费不含紧急购电，实际总费用另行列示。问题3采用策略D；问题4的明细对应当日价格曲线已知的条件基准，历史价格预测方案汇总见5.5节。')
def group_tables(tag,label):
    h(label,3)
    recs=[r for r in typ if r['model']==tag]
    for r in recs:
        date=r['date'];s=r['slots'];b=r['blocks']
        tab(date+'日前购电量及全天结果',['时间段','购电量','时间段','购电量','时间段','购电量'],[[s[j]['time'],f(s[j]['plan'],4),s[j+1]['time'],f(s[j+1]['plan'],4),s[j+2]['time'],f(s[j+2]['plan'],4)] for j in [0,3]]+[['全天购电量',f(r['G0'],4),'全天购电费',f(r['base']),'','']],[2.8,2.53,2.8,2.53,2.8,2.54],10)
        tab(date+'实际储能充放电量',['时间段','充电量','放电量','时间段','充电量','放电量'],[[b[j]['time'],f(b[j]['charge'],4),f(b[j]['discharge'],4),b[j+1]['time'],f(b[j+1]['charge'],4),f(b[j+1]['discharge'],4)] for j in [0,2,4]]+[['0:00储电量',f(r['initial'],4),'','24:00储电量',f(r['final'],4),'']],[2.8,2.6,2.6,2.8,2.6,2.6],10)
    heads=[]
    for r in recs:heads.extend([r['date'][5:]+'时间段',r['date'][5:]+'购电量'])
    data=[]
    for j in range(max(len(r['events']) for r in recs)):
        line=[]
        for r in recs:line.extend([r['events'][j][0],f(r['events'][j][1],4)] if j<len(r['events']) else ['—','—'])
        data.append(line)
    tab('四个指定日期的紧急购电明细',heads,data,[2.2,1.8]*4,9)
    tab('指定日期的完整费用核算',['日期','计划费','净调整费','紧急费','总费用'],[[r['date'],f(r['base']),f(r['adjustment']),f(r['emergency']),f(r['total'])] for r in recs],[3,3.2,3.2,3.2,3.4],10)
    if tag in ['q3_D','q43_D']:
        tab('调整后的最终常规购电量',['日期','10:00','12:00','14:00','16:00','18:00','20:00','全天'],[[r['date'][5:]]+[f(s['final'],2) for s in r['slots']]+[f(r['G'],2)] for r in recs],[1.4,1.8,1.8,1.8,1.8,1.8,1.8,3.8],9.5)
        p('上表各整点列表示该时刻开始的10分钟购电量，展示精度为两位小数，完整精度见结果表。“最终常规购电”已包含各次调整，不能再与增购量重复相加。')
    p('同一四小时区间内充、放电累计量均为正，表示不同10分钟时段之间发生运行切换，并不表示同一时刻同时充放电。紧急区间由相邻正值时段合并，空余位置以横线表示。')
for tag,label in [('q2','5.6.1 问题2的指定日期结果'),('q3_D','5.6.2 问题3的指定日期结果'),('q42','5.6.3 问题4-2的指定日期结果'),('q43_D','5.6.4 问题4-3的指定日期结果')]:group_tables(tag,label)

h('5.7 模型检验与敏感性分析',2)
h('5.7.1 物理约束与费用一致性检验',3)
p(r'对12组全年策略分别检查11项指标，共132项检验，最大绝对误差为$5.82\times10^{-11}$，低于$10^{-6}$的数值容差。检验覆盖能量守恒、储能状态、容量、功率、非负性、充放电互斥、跨日连续、费用、交易恒等式、紧急区间及有限数值。')
p(r'独立读取五份结果工作簿，与重建轨迹比较288720个时段购电值，并检查储能汇总、日期与费用，最大差异约$1.2\times10^{-10}$。另进行61项因果污染测试，检查改变未来实际值或后续预报不会影响此前计划。测试支持已检查路径的信息边界，但不替代对全部控制逻辑的数学证明。')
h('5.7.2 关键参数敏感性分析',3)
p('选取3月18—23日、6月19—24日、9月21—26日及12月19—24日，共24天进行对比。各组采用相同初始储电量，并用库存价值校正终端状态差异：')
eq(r'J_{\mathrm{adj}}=J-0.5\overline p\,(S_{\mathrm{end}}-S_{\mathrm{start}})')
ss={r['case']:r for r in rows('敏感性24日.csv') if r['model']=='q3_D'};base=float(ss['baseline']['inventory_adjusted_cost'])
tab('问题3策略D的参数敏感性',['参数情形','库存校正费用 / 元','相对基准变化'],[[lab,f(ss[key]['inventory_adjusted_cost']),f'{(float(ss[key]["inventory_adjusted_cost"])/base-1)*100:+.2f}%'] for key,lab in [('baseline','基准参数'),('eta085','充放电效率各85%'),('eta095','充放电效率各95%'),('roundtrip090','往返效率90%'),('capacity9600','容量9600 kWh'),('capacity14400','容量14400 kWh'),('margin1','残差裕度1'),('no_refund','取消不退款'),('pchip','PCHIP插值')]],[6.3,5.5,4.2])
p('在所测范围内，提高效率与扩大容量能够降低费用；增加安全裕度可能因额外购电和弃置而提高成本。取消不退款时，策略D费用上升，但未在该条款下重新比较全部A—D组合，因此不能断言所有更新时点仍然有利。24天结果用于说明局部变化趋势，不外推为全年收益保证。')

fig(HERE/'新增图/06_敏感性对比.png','关键参数变化对24日库存校正费用的影响',15.5)
p('以零变化线为参照，左侧表示费用下降，右侧表示费用上升。该图与敏感性表使用相同的策略D和24日样本，便于比较效率、容量及结算参数的影响方向，不能据此代替全年重新求解。')
h('第6章　模型评价与改进')
h('6.1 模型的优点',2)
p('本文以统一的能量平衡和储能状态连接四个问题，使功率、电量、库存与费用之间的关系清晰。确定性问题可通过线性规划得到最优解，并给出对偶检验；不确定性问题遵循信息发布时间进行连续回放，计划、调整和紧急购电分别计量，能够追溯指定日期结果。')
p('模型将预测方法的经济选型与外样本评价分开，并同时报告波动电价的两种信息条件，有助于避免将事后已知数据误用于决策。对容量、效率及结算条款的敏感性检验，也使结论的适用范围更加明确。')
h('6.2 模型的不足',2)
p('三条历史残差情景不能完整刻画极端天气和尾部风险，情景内动作缺少完整的多阶段非预见性约束。实际电池采用贪心执行规则，没有显式优化未来储能机会成本。因此，问题2至问题4所得费用属于所构造策略的回放结果，不具有全局最优保证。')
p('模型未考虑自放电、寿命损耗、售电和外网交易功率限制，结算规则也采用了明确的退款约定。这些简化与题目数据范围相适应，但推广到实际微网前需要补充设备与市场约束。')
h('6.3 模型的改进方向',2)
p('可将电池执行规则改进为基于条件预测的短时域优化，并在场景树中加入非预见性约束，使储能决策兼顾后续风险。对于新增价格发布机制、退款条款或预报时点，应重新训练和回放，通过统一信息边界、终端价值与实际账单比较评估收益。')
p(f'综合结果，问题1的最低日购电费为{f(cert["cost"])}元；问题2与问题3策略D的2月至12月费用分别为{f(val("q2","total")/10000)}万元和{f(val("q3_D","total")/10000)}万元。已有三个额外预报时点在所测嵌套策略中均带来节省，波动价格下的结论则需结合价格信息条件解释。')
h('参考文献')
refs=[
'[1] SciPy Developers. scipy.optimize.linprog[EB/OL]. https://docs.scipy.org/doc/scipy/reference/generated/scipy.optimize.linprog.html, 2026-09-10.',
'[2] Boyd S, Vandenberghe L. Convex Optimization[M]. Cambridge: Cambridge University Press, 2004.',
'[3] Hyndman R J, Athanasopoulos G. Forecasting: Principles and Practice[M/OL]. 3rd ed. Melbourne: OTexts, 2021. https://otexts.com/fpp3/.',
'[4] Rawlings J B, Mayne D Q, Diehl M M. Model Predictive Control: Theory, Computation, and Design[M]. 2nd ed., 6th printing. Nob Hill Publishing, 2026.'
]
for text in refs:
    pp=p(text,False);pp.alignment=WD_ALIGN_PARAGRAPH.LEFT;pp.paragraph_format.line_spacing=Pt(18);pp.paragraph_format.space_after=Pt(5)
    for rr in pp.runs:rr.font.size=Pt(10.5)

br();h('附录A　支撑材料与运行说明')
p('支撑材料位于工程的审查与完善目录。input保存题设附件，results保存逐时段轨迹和逐日费用，audit保存验证结果，figures保存图形。下表列出主要计算文件；完整计算程序随后附出，排版与公式转换代码单独保存在本次修订目录及原复现目录中。')
tab('支撑材料文件列表',['文件或目录','用途'],[['src/model.py','数据解析、预测、调度及物理与费用检查'],['src/run.py','四问计算、参数选型及敏感性实验'],['src/q1_certificate.py','独立线性规划及最优性检验'],['src/review_results.py','结果表独立回读与指定日期汇总'],['src/deliver.py中的causality函数','未来输入污染测试'],['src/figures.py','原有计算结果图形生成'],['make_added_figures.py','本次新增6幅图及数据溯源'],['results/*.npz；*_daily.csv','逐时段轨迹及逐日账单'],['audit/*.csv；*.json','数据与约束检查'],['06_错误说明与纠正/问题1至问题4','五份修订结果工作簿']],[7,9])
p('运行环境为Python 3.12及NumPy、SciPy、pandas、openpyxl、matplotlib等依赖。先在原复现目录执行run.py的all阶段，再执行review_results.py与独立验证。输入附件放入input；q1_certificate.py按其文件位置读取工程附件，运行前应保持目录结构。具体命令与依赖版本见既有运行说明。')
p('本次使用OpenAI Codex辅助代码审查、计算复核、文字修订和公式排版。数值来自题设附件与所列程序。该说明不替代正式参赛所需的完整AI工具交互记录；参赛者应据实际使用情况整理相关材料。')
for k,fname in enumerate(['model.py','run.py','review_results.py','q1_certificate.py','deliver.py','figures.py','make_added_figures.py']):
    br();h(f'附录{chr(66+k)}　{fname}源程序')
    source=((HERE if fname=='make_added_figures.py' else SRC)/fname).read_text(encoding='utf-8')
    if fname=='deliver.py':
        p('本程序保留完整上下文以便复现因果检验。复现检验仅调用causality函数；export函数为历史导出实现，不用于本次修订表的生成。')
    for line in source.splitlines():
        pp=doc.add_paragraph();pf=pp.paragraph_format;pf.first_line_indent=Pt(0);pf.space_before=pf.space_after=Pt(0);pf.line_spacing=Pt(10);pf.alignment=WD_ALIGN_PARAGRAPH.LEFT;pf.widow_control=False
        rr=pp.add_run(line or ' ');rr.font.name='Consolas';rr.font.size=Pt(8);rr._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
for el in list(doc.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
for pp in doc.paragraphs:
    if pp.text=='表5-51 支撑材料文件列表':pp.runs[0].text='表A-1 支撑材料文件列表'
for run in doc.element.iter(qn('m:r')):
    text_node=run.find(qn('m:t'))
    if text_node is not None and text_node.text in ['min','max','exp']:
        pr=run.find(qn('m:rPr'))
        if pr is None:pr=OxmlElement('m:rPr');run.insert(0,pr)
        sty=pr.find(qn('m:sty'))
        if sty is None:sty=OxmlElement('m:sty');pr.append(sty)
        sty.set(qn('m:val'),'p')
doc.core_properties.title='基于滚动预测与储能协同调度的微网购电策略优化研究'
OUT=HERE/'微网购电策略论文_图文增强版.docx';doc.save(OUT)
(HERE/'论文正文与公式源.md').write_text('\n\n'.join(notes),encoding='utf-8')
(HERE/'公式LaTeX源.tex').write_text('\n\n'.join('\\[\n'+x+'\n\\]' for x in FORMULAS),encoding='utf-8')
print(json.dumps({'output':str(OUT),'formulas':len(FORMULAS),'display_formulas':n_eq,'tables':len(doc.tables)},ensure_ascii=False))
